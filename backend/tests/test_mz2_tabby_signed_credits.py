"""Actual signed statement credits: no contract-rate inference or customer reversal."""
import io
import os
import unittest
from uuid import uuid4

import openpyxl
from fastapi import HTTPException
from accounting_settlement_service import amounts_from_settlement_file, build_journal_preview, post_reviewed_settlement
from settlements_import.parsers.tabby import parse


def statement():
    wb = openpyxl.Workbook()
    wb.active.append(["Statement #", "SYN-SIGNED-CREDITS"])
    wb.active.append(["Date", "2026-09-07"])
    wb.active.append(["Order Number", "Sale/Refund Date", "Type", "Order Amount",
        "Refundable Commission", "Total Fee", "VAT Amount", "Transferred amount"])
    wb.active.append([9911030998, "2026-09-02", "sale", 200, 1, 1, .15, 198.85])
    wb.active.append([9911030999, "2026-08-25", "refund", 115, -3.21, -3.21, -.48, -111.31])
    return wb


class SignedCreditTests(unittest.TestCase):
    def test_prior_period_document_keeps_signed_components_and_balances(self):
        parsed = parse(statement())
        self.assertEqual(parsed["entries"][1]["event_date"], "2026-08-25")
        self.assertEqual(parsed["totals"]["refunded_fees"], -3.21)
        self.assertEqual(parsed["totals"]["refunded_fees_vat"], -.48)
        amounts = amounts_from_settlement_file(parsed)
        self.assertEqual((amounts["commission"], amounts["commission_vat"]), (-2.21, -.33))
        preview = build_journal_preview(provider="tabby", bank_account_id="b", bank_account_name="Test", amounts=amounts)
        self.assertTrue(preview["balanced"])
        self.assertEqual(preview["debit_total"], 88.69)
        self.assertEqual(preview["credit_total"], 88.69)
        legs = {e["role"]: e for e in preview["entries"]}
        self.assertEqual((legs["commission"]["side"], legs["commission"]["amount"]), ("debit", 1))
        self.assertEqual((legs["commission_credit"]["side"], legs["commission_credit"]["source_signed_amount"]), ("credit", -3.21))
        self.assertEqual(legs["commission_vat_credit"]["source_signed_amount"], -.48)
        self.assertEqual(legs["provider_receivable"]["amount"], 85)
        self.assertFalse(any(e["entity_type"] in {"tax", "revenue", "liability"} for e in preview["entries"]))

    def test_positive_refund_fee_is_a_charge_not_an_automatic_rebate(self):
        wb = statement()
        wb.active.cell(5, 6, 3.21)
        wb.active.cell(5, 7, .48)
        wb.active.cell(5, 8, -118.69)
        parsed = parse(wb)
        self.assertEqual(parsed["totals"]["fee_credits"], 0)
        self.assertEqual(parsed["totals"]["refunded_fees"], 3.21)
        preview = build_journal_preview(provider="tabby", bank_account_id="b", bank_account_name="Test", amounts=amounts_from_settlement_file(parsed))
        self.assertTrue(preview["balanced"])
        self.assertFalse(any(e["side"] == "credit" and e["entity_type"] == "expense" for e in preview["entries"]))

    def test_other_provider_cannot_use_signed_tabby_credit(self):
        for provider in ("salla", "tamara", "emkan"):
            with self.assertRaisesRegex(ValueError, "documented_tabby"):
                build_journal_preview(provider=provider, bank_account_id="b", bank_account_name="Test", amounts=amounts_from_settlement_file(parse(statement())))


@unittest.skipUnless(os.getenv("MZ2_TEST_MONGO_URI"), "dedicated replica URI required")
class SignedCreditMongoTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        from motor.motor_asyncio import AsyncIOMotorClient
        from ledger_core import post_txn_group
        from accounting_atomic import atomic_owner
        from accounting_module_contract import OPERATION_ID
        self.mongo = AsyncIOMotorClient(os.environ["MZ2_TEST_MONGO_URI"])
        self.db = self.mongo["mz2_atomic_test_" + uuid4().hex]
        self.assertTrue((await self.db.command("hello")).get("setName"))
        await self.db.users.insert_one({"id": "owner", "role": "owner"})
        await self.db.settings.insert_one({"user_id": "owner", "mezan2_financial_cutover": {
            "operation_id": OPERATION_ID, "cutover_at": "2020-01-01T00:00:00Z"}})
        await self.db.accounts.insert_one({"user_id": "owner", "id": "b", "name": "Synthetic bank", "account_type": "bank"})
        async def opening(scoped):
            return await post_txn_group(scoped, user_id="owner", actor_id="owner", actor_name="Test",
                entries=[dict(entity_type="payment_gateway", entity_id="tabby", sub_account="receivable", side="debit", amount=85, entry_type="opening_balance"),
                         dict(entity_type="revenue", entity_id="fixture", side="credit", amount=85, entry_type="opening_balance")],
                txn_type="synthetic_opening", reason_code="fixture", metadata={"accounting_at": "2026-08-01T00:00:00Z"})
        opening_group = await atomic_owner(self.db, "owner", opening)
        from mz2_report_fixtures import provision_write_opening
        await provision_write_opening(self.db, existing_group_id=opening_group['txn_group_id'],
            bank_zero_ids=('b',), providers=('salla', 'tamara', 'emkan'))
        from settlements_import.service import import_file
        out = io.BytesIO(); statement().save(out); self.content = out.getvalue()
        async def upload(scoped):
            return await import_file(scoped, user_id="owner", content=self.content, filename="synthetic.xlsx", provider_hint="tabby")
        result = await atomic_owner(self.db, "owner", upload)
        self.file = await self.db.settlement_files.find_one({"id": result["file_id"]})
        row = await self.db.settlement_entries.find_one({"file_id": self.file["id"], "event_type": "refund"})
        await self.db.mz2_statement_refund_links.insert_one({"user_id": "owner", "draft_id": "d", "entry_id": row["id"], "amount": "115"})
        self.draft = dict(id="d", status="reviewed", provider="tabby", bank_account_id="b", idempotency_key="signed-credit",
            statement_reference="SYN-SIGNED-CREDITS", statement_date="2026-09-07", source_file_id=self.file["id"],
            source_file_hash=self.file["file_hash"], amounts=amounts_from_settlement_file(self.file))

    async def asyncTearDown(self):
        self.mongo.close()

    async def test_real_post_renamed_upload_duplicate_and_source_guard(self):
        from settlements_import.service import import_file
        from accounting_atomic import atomic_owner
        result = await post_reviewed_settlement(self.db, owner_id="owner", actor={"id": "owner"}, draft=self.draft)
        legs = await self.db.general_ledger.find({"txn_group_id": result["txn_group_id"]}).to_list(20)
        self.assertEqual(len(legs), 6)
        self.assertEqual(legs[0]["metadata"]["accounting_at"], "2026-09-06T21:00:00+00:00")
        self.assertEqual(legs[0]["metadata"]["fee_credit_evidence"][0]["actual_payment_fee"], -3.21)
        before = await self.db.general_ledger.count_documents({})
        async def upload(scoped):
            return await import_file(scoped, user_id="owner", content=self.content, filename="renamed.xlsx", provider_hint="tabby")
        duplicate = await atomic_owner(self.db, "owner", upload)
        self.assertEqual(duplicate["status"], "duplicate")
        self.assertEqual(duplicate["file_id"], self.file["id"])
        with self.assertRaises(HTTPException) as repeated:
            await post_reviewed_settlement(self.db, owner_id="owner", actor={"id": "owner"}, draft=self.draft)
        self.assertEqual(repeated.exception.status_code, 409)
        self.assertEqual(await self.db.general_ledger.count_documents({}), before)
        changed = {**self.draft, "idempotency_key": "changed", "amounts": {**self.draft["amounts"], "commission_credit": -5}}
        with self.assertRaises(HTTPException) as invalid:
            await post_reviewed_settlement(self.db, owner_id="owner", actor={"id": "owner"}, draft=changed)
        self.assertEqual(invalid.exception.detail, "tabby_fee_credit_source_conflict")
        self.assertEqual(await self.db.general_ledger.count_documents({}), before)
