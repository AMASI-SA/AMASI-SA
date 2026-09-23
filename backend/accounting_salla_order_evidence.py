"""MZ2-native Salla order-export evidence intake.

This module is intentionally evidence-first:
- Uploading a Salla order workbook never posts a journal.
- Customer PII is not copied into the normalized accounting evidence rows.
- The immutable original workbook is preserved separately for audit.
- Payment identities come only from Salla's explicit payment-reference JSON.
- COD and bank-transfer rows are routed to their later evidence workflows.
- Amount/refund/status conflicts become needs-review states; they are never guessed.
"""
from __future__ import annotations

from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
import hashlib
import io
import json
import re
import uuid
from typing import Any

from fastapi import Depends, File, HTTPException, UploadFile
from openpyxl import load_workbook

from accounting_atomic import atomic_owner
from accounting_module_contract import accounting_owner_id, require_accounting_permission
from accounting_module_status_routes import fresh_accounting_user
from accounting_source_files import preserve_original
from excel_upload_security import read_safe_xlsx_upload


MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_ROWS = 50_000
MONEY = Decimal("0.01")

REQUIRED_HEADERS = {
    "رقم الطلب",
    "حالة الطلب",
    "طريقة الدفع",
    "رقم مرجع عملية الدفع",
    "صافي المبيعات",
    "تاريخ الطلب",
    "تاريخ آخر تحديث للطلب",
    "إجمالي الطلب بالعملة الأصلية",
    "العملة الأصلية للطلب",
    "المبلغ المسترجع",
}
OPTIONAL_HEADERS = {
    "تفاصيل طرق الدفع",
    "تاريخ التسليم",
    "شركة الشحن / الفرع",
    "شركة الشحن",
    "تكلفة الشحن",
    "عمولة الدفع عند الاستلام",
    "رقم البوليصة",
    "رقم مرجع الطلب",
    "إجمالي المبيعات",
    "عملة الطلب",
    "الضريبة",
}

SALLA_METHODS = {
    "مدى", "البطاقه الائتمانيه", "البطاقة الائتمانية", "البطاقة الإئتمانية",
    "stc pay", "stcpay",
}
TAMARA_METHODS = {"تمارا", "tamara"}
TABBY_METHODS = {"تابي", "tabby"}
EMKAN_METHODS = {"emkaninstallment", "emkan", "imkan", "إمكان", "امكان"}
COD_METHODS = {"دفع عند الاستلام", "دفع عند الإستلام", "الدفع عند الاستلام", "cod", "cash on delivery"}
BANK_TRANSFER_HINTS = {"حوالة بنكية", "تحويل بنكي", "bank transfer"}

PAYMENT_PROVIDER_ALIASES = {
    "salla": {"applepay", "checkout", "stcpay", "mada", "visa", "mastercard", "salla", "salla_pay"},
    "tamara": {"tamara"},
    "tabby": {"tabby"},
    "emkan": {"emkan", "imkan"},
}

READY_STATUS = "تم التوصيل"


def _clean(value: Any) -> str:
    if value is None:
        return ""
    return " ".join(str(value).strip().strip("'").split())


def _norm(value: Any) -> str:
    text = _clean(value).lower()
    text = re.sub(r"[\u064B-\u0652]", "", text)
    return (
        text.replace("أ", "ا").replace("إ", "ا").replace("آ", "ا")
        .replace("ى", "ي").replace("ة", "ه")
    )


def _money(value: Any) -> Decimal:
    if value in (None, ""):
        return Decimal("0.00")
    try:
        result = Decimal(str(value)).quantize(MONEY, rounding=ROUND_HALF_UP)
    except (InvalidOperation, ValueError, TypeError):
        raise ValueError("invalid_money") from None
    if not result.is_finite() or result < 0:
        raise ValueError("invalid_money")
    return result


def _timestamp(value: Any) -> str | None:
    text = _clean(value)
    if not text:
        return None
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d"):
        try:
            parsed = datetime.strptime(text, fmt).replace(tzinfo=timezone.utc)
            # Source export carries no timezone marker. Store source text too;
            # this normalized value is ordering-only and never becomes an
            # accounting instant by itself.
            return parsed.isoformat()
        except ValueError:
            pass
    return None


def _hash(value: Any) -> str:
    raw = json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str)
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()


def _payment_method(raw: Any) -> str:
    value = _norm(raw)
    if value in {_norm(v) for v in SALLA_METHODS}:
        return "salla"
    if value in {_norm(v) for v in TAMARA_METHODS}:
        return "tamara"
    if value in {_norm(v) for v in TABBY_METHODS}:
        return "tabby"
    if value in {_norm(v) for v in EMKAN_METHODS}:
        return "emkan"
    if value in {_norm(v) for v in COD_METHODS}:
        return "cod"
    if any(_norm(hint) in value for hint in BANK_TRANSFER_HINTS):
        return "bank_transfer"
    return "unknown"


def _payment_reference(value: Any) -> dict[str, Any] | None:
    text = _clean(value)
    if not text or text in {"\\N", "null", "None"}:
        return None
    try:
        payload = json.loads(text)
    except json.JSONDecodeError:
        raise ValueError("payment_reference_json_invalid") from None
    if not isinstance(payload, list) or len(payload) != 1 or not isinstance(payload[0], dict):
        raise ValueError("payment_reference_shape_unsupported")
    item = payload[0]
    reference = _clean(item.get("reference"))
    provider = _clean(item.get("provider"))
    amount = _money(item.get("amount"))
    if not reference or not provider or amount <= 0:
        raise ValueError("payment_reference_incomplete")
    return {
        "provider_raw": provider,
        "provider_normalized": _norm(provider),
        "reference": reference,
        "amount": format(amount, ".2f"),
    }


def _provider_reference_matches(accounting_provider: str, payment_ref: dict[str, Any] | None) -> bool:
    if not payment_ref or accounting_provider not in PAYMENT_PROVIDER_ALIASES:
        return False
    return payment_ref["provider_normalized"] in {
        _norm(alias) for alias in PAYMENT_PROVIDER_ALIASES[accounting_provider]
    }


def _headers(sheet) -> tuple[int, list[str], dict[str, int]]:
    for row_no, row in enumerate(sheet.iter_rows(min_row=1, max_row=20, values_only=True), start=1):
        values = [_clean(v) for v in row]
        present = {v for v in values if v}
        if "رقم الطلب" in present and "طريقة الدفع" in present:
            missing = sorted(REQUIRED_HEADERS - present)
            if missing:
                raise ValueError("missing_required_columns:" + "|".join(missing))
            mapping = {name: values.index(name) for name in present if name in REQUIRED_HEADERS | OPTIONAL_HEADERS}
            return row_no, values, mapping
    raise ValueError("salla_order_header_not_detected")


def _cell(row: tuple[Any, ...], mapping: dict[str, int], name: str) -> Any:
    index = mapping.get(name)
    return row[index] if index is not None and index < len(row) else None


def classify_salla_order_row(row: dict[str, Any]) -> tuple[str, list[str]]:
    reasons: list[str] = []
    provider = row["accounting_provider"]
    status = row["order_status"]
    delivered_at = row.get("delivery_source_text")
    payment_ref = row.get("payment_reference")
    current_sar = Decimal(row["current_net_sar"])
    refunded = Decimal(row["refunded_sar"])

    if provider == "unknown":
        return "needs_review", ["payment_method_unknown"]
    if provider == "cod":
        return "waiting_p02_cod", []
    if provider == "bank_transfer":
        return "needs_bank_transfer_evidence", []

    if payment_ref is None:
        reasons.append("payment_reference_missing")
    elif not _provider_reference_matches(provider, payment_ref):
        reasons.append("payment_reference_provider_conflict")
    else:
        paid = Decimal(payment_ref["amount"])
        if refunded > 0:
            if paid < refunded or (paid - refunded) != current_sar:
                reasons.append("refund_amount_conflict")
        elif paid != current_sar:
            reasons.append("payment_amount_conflict")

    if status != READY_STATUS or not delivered_at:
        reasons.append("fulfilment_timestamp_required")

    if reasons:
        return "needs_review", sorted(set(reasons))
    if refunded > 0:
        return "ready_sale_refund_pending_evidence", ["refund_provider_identity_required"]
    return "ready_for_provider_resolution", []


def parse_salla_order_xlsx(content: bytes) -> dict[str, Any]:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
    try:
        sheet = workbook.active
        header_row, _header_values, mapping = _headers(sheet)
        rows: list[dict[str, Any]] = []
        errors: list[dict[str, Any]] = []
        seen_orders: set[str] = set()

        for row_no, raw in enumerate(
            sheet.iter_rows(min_row=header_row + 1, values_only=True),
            start=header_row + 1,
        ):
            if not any(value not in (None, "") for value in raw):
                continue
            if len(rows) + len(errors) >= MAX_ROWS:
                raise ValueError("salla_order_row_limit")
            try:
                order_number = _clean(_cell(raw, mapping, "رقم الطلب"))
                if not order_number:
                    raise ValueError("order_number_required")
                if order_number in seen_orders:
                    raise ValueError("duplicate_order_number_in_file")
                seen_orders.add(order_number)

                payment_method_raw = _clean(_cell(raw, mapping, "طريقة الدفع"))
                provider = _payment_method(payment_method_raw)
                payment_ref = _payment_reference(_cell(raw, mapping, "رقم مرجع عملية الدفع"))

                original_currency = _clean(_cell(raw, mapping, "العملة الأصلية للطلب")).upper() or "SAR"
                original_amount = _money(_cell(raw, mapping, "إجمالي الطلب بالعملة الأصلية"))
                current_sar = _money(_cell(raw, mapping, "صافي المبيعات"))
                refunded = _money(_cell(raw, mapping, "المبلغ المسترجع"))
                shipping_cost = _money(_cell(raw, mapping, "تكلفة الشحن"))
                cod_fee = _money(_cell(raw, mapping, "عمولة الدفع عند الاستلام"))

                if original_amount <= 0 or current_sar < 0:
                    raise ValueError("order_amount_invalid")

                order_status = _clean(_cell(raw, mapping, "حالة الطلب"))
                order_date_text = _clean(_cell(raw, mapping, "تاريخ الطلب"))
                updated_text = _clean(_cell(raw, mapping, "تاريخ آخر تحديث للطلب"))
                delivery_text = _clean(_cell(raw, mapping, "تاريخ التسليم"))
                if not order_date_text or not updated_text:
                    raise ValueError("order_source_date_required")

                row = {
                    "row_no": row_no,
                    "order_number": order_number,
                    "order_reference": _clean(_cell(raw, mapping, "رقم مرجع الطلب")) or None,
                    "order_status": order_status,
                    "payment_method_raw": payment_method_raw,
                    "accounting_provider": provider,
                    "payment_reference": payment_ref,
                    "current_net_sar": format(current_sar, ".2f"),
                    "refunded_sar": format(refunded, ".2f"),
                    "original_amount": format(original_amount, ".2f"),
                    "original_currency": original_currency,
                    "order_date_source_text": order_date_text,
                    "order_date_ordering": _timestamp(order_date_text),
                    "updated_source_text": updated_text,
                    "updated_ordering": _timestamp(updated_text),
                    "delivery_source_text": delivery_text or None,
                    "delivery_ordering": _timestamp(delivery_text),
                    "shipping_company": _clean(
                        _cell(raw, mapping, "شركة الشحن / الفرع")
                        or _cell(raw, mapping, "شركة الشحن")
                    ) or None,
                    "shipping_cost_source": format(shipping_cost, ".2f"),
                    "cod_fee_source": format(cod_fee, ".2f"),
                    "source_tax_sar": format(_money(_cell(raw, mapping, "الضريبة")), ".2f"),
                    "waybill": _clean(_cell(raw, mapping, "رقم البوليصة")) or None,
                }
                state, reasons = classify_salla_order_row(row)
                row["status"] = state
                row["review_reasons"] = reasons
                row["economic_hash"] = _hash({
                    key: row[key] for key in (
                        "order_number", "order_status", "payment_method_raw",
                        "accounting_provider", "payment_reference", "current_net_sar",
                        "refunded_sar", "original_amount", "original_currency",
                        "order_date_source_text", "updated_source_text",
                        "delivery_source_text", "shipping_company",
                        "shipping_cost_source", "cod_fee_source", "source_tax_sar", "waybill",
                    )
                })
                rows.append(row)
            except ValueError as exc:
                errors.append({
                    "row_no": row_no,
                    "order_number": _clean(_cell(raw, mapping, "رقم الطلب")) or None,
                    "code": str(exc),
                })

        if not rows and not errors:
            raise ValueError("salla_order_file_has_no_rows")
        return {
            "header_row": header_row,
            "rows": rows,
            "errors": errors,
        }
    finally:
        workbook.close()


RECOGNITION_FIELDS = (
    "recognition_event_key",
    "recognition_txn_group_id",
    "recognized_at",
    "recognized_provider",
    "recognized_gross_sar",
    "recognized_tax_sar",
    "recognized_net_sar",
    "recognized_by",
    "recognition_posted_at",
)


BANK_TRANSFER_FIELDS = (
    "bank_selected_from_order",
    "bank_transfer_receipt_review_id",
    "bank_transfer_receipt_status",
    "bank_transfer_advance_id",
    "bank_transfer_receipt_txn_group_id",
    "bank_transfer_sale_txn_group_id",
)


def _preserve_bank_transfer_state(
    prior: dict[str, Any],
    current: dict[str, Any],
) -> tuple[dict[str, Any], str | None]:
    if not prior.get("bank_transfer_receipt_review_id"):
        return current, None
    for field in BANK_TRANSFER_FIELDS:
        if field in prior:
            current[field] = prior[field]
    if current.get("conflict"):
        return current, None

    reason = None
    if current.get("accounting_provider") != "bank_transfer":
        reason = "bank_transfer_payment_method_changed"
    elif str(prior.get("current_net_sar") or "") != str(current.get("current_net_sar") or ""):
        reason = "bank_transfer_order_amount_changed"
    elif Decimal(str(current.get("refunded_sar") or "0")) != Decimal("0"):
        reason = "bank_transfer_refund_requires_review"
    else:
        state = prior.get("bank_transfer_receipt_status")
        if state == "recognized":
            current["status"] = "recognized_bank_transfer"
            current["review_reasons"] = []
        elif state == "confirmed_waiting_delivery":
            current["status"] = "bank_transfer_confirmed_waiting_delivery"
            current["review_reasons"] = []
        elif state == "pending_approval":
            current["status"] = "needs_bank_transfer_evidence"
            current["review_reasons"] = []

    if reason:
        current["conflict"] = True
        current["status"] = "needs_review"
        current["review_reasons"] = sorted(set(
            list(current.get("review_reasons") or []) + [reason]
        ))
    return current, reason


def _preserve_recognized_state(prior: dict[str, Any], current: dict[str, Any]) -> tuple[dict[str, Any], str | None]:
    """Keep a posted sale attached when a newer Salla export arrives.

    Shipping/display changes are allowed.  A later, reconciled refund moves the
    row to refund-pending.  Changes to the already-posted payment identity,
    gross amount, or delivery accounting source date become review conflicts
    instead of silently rewriting history.
    """
    if not prior.get("recognition_txn_group_id"):
        return current, None
    for field in RECOGNITION_FIELDS:
        if field in prior:
            current[field] = prior[field]
    if current.get("conflict"):
        return current, None

    reason = None
    old_payment = prior.get("payment_reference") or {}
    new_payment = current.get("payment_reference") or {}
    if (
        prior.get("recognized_provider") != current.get("accounting_provider")
        or old_payment.get("reference") != new_payment.get("reference")
        or str(old_payment.get("amount") or "") != str(new_payment.get("amount") or "")
    ):
        reason = "recognized_payment_identity_changed"
    elif prior.get("delivery_source_text") != current.get("delivery_source_text"):
        reason = "recognized_delivery_timestamp_changed"
    else:
        gross = Decimal(str(prior.get("recognized_gross_sar") or "0"))
        refunded = Decimal(str(current.get("refunded_sar") or "0"))
        current_net = Decimal(str(current.get("current_net_sar") or "0"))
        if refunded > 0 and gross - refunded == current_net:
            current["status"] = "recognized_refund_pending_evidence"
            current["review_reasons"] = ["refund_provider_identity_required"]
        elif refunded == 0 and gross == current_net:
            current["status"] = "recognized"
            current["review_reasons"] = []
        else:
            reason = "recognized_amount_changed"

    if reason:
        current["conflict"] = True
        current["status"] = "needs_review"
        current["review_reasons"] = sorted(set(
            list(current.get("review_reasons") or []) + [reason]
        ))
    return current, reason


async def import_salla_order_evidence(
    db,
    *,
    owner: str,
    actor: dict[str, Any],
    filename: str,
    content: bytes,
    parsed: dict[str, Any],
) -> dict[str, Any]:
    file_hash = hashlib.sha256(content).hexdigest()
    file_key = _hash([owner, "salla_order_export", file_hash])
    existing_file = await db.mz2_salla_order_files.find_one({"_id": file_key, "user_id": owner})
    if existing_file:
        items = await db.mz2_salla_order_snapshots.find(
            {"user_id": owner, "file_id": existing_file["id"]},
            {"_id": 0},
        ).sort("row_no", 1).to_list(MAX_ROWS)
        return {
            "status": "duplicate",
            "file": {k: v for k, v in existing_file.items() if k != "_id"},
            "items": items,
            "errors": existing_file.get("errors") or [],
        }

    file_id = str(uuid.uuid4())
    now = datetime.now(timezone.utc).isoformat()
    await preserve_original(db, owner, file_id, content)

    snapshots = []
    current_updates = []
    conflicts = []
    counts: dict[str, int] = {}

    for row in parsed["rows"]:
        snapshot_id = _hash([owner, row["order_number"], file_hash])
        snapshot = {
            "_id": snapshot_id,
            "id": snapshot_id,
            "user_id": owner,
            "file_id": file_id,
            "file_hash": file_hash,
            "source": "salla_orders_export",
            "imported_at": now,
            **row,
        }
        snapshots.append(snapshot)
        counts[row["status"]] = counts.get(row["status"], 0) + 1

        current_id = _hash([owner, "salla_order", row["order_number"]])
        prior = await db.mz2_salla_order_evidence.find_one({"_id": current_id, "user_id": owner})
        conflict = False
        if prior:
            prior_updated = str(prior.get("updated_source_text") or "")
            incoming_updated = str(row.get("updated_source_text") or "")
            if incoming_updated < prior_updated:
                continue
            if incoming_updated == prior_updated and prior.get("economic_hash") != row["economic_hash"]:
                conflict = True
                conflicts.append({
                    "order_number": row["order_number"],
                    "code": "same_source_version_changed",
                    "prior_snapshot_id": prior.get("latest_snapshot_id"),
                    "incoming_snapshot_id": snapshot_id,
                })
        current = {
            "_id": current_id,
            "id": current_id,
            "user_id": owner,
            "source": "salla_orders_export",
            "latest_snapshot_id": snapshot_id,
            "latest_file_id": file_id,
            "latest_file_hash": file_hash,
            "conflict": conflict,
            **row,
        }
        if conflict:
            current["status"] = "needs_review"
            current["review_reasons"] = sorted(set(
                list(current.get("review_reasons") or []) + ["same_source_version_changed"]
            ))
        if prior and prior.get("recognition_txn_group_id"):
            current, recognized_conflict = _preserve_recognized_state(prior, current)
            if recognized_conflict:
                conflicts.append({
                    "order_number": row["order_number"],
                    "code": recognized_conflict,
                    "prior_snapshot_id": prior.get("latest_snapshot_id"),
                    "incoming_snapshot_id": snapshot_id,
                    "recognition_txn_group_id": prior.get("recognition_txn_group_id"),
                })
        if prior and prior.get("bank_transfer_receipt_review_id"):
            current, bank_conflict = _preserve_bank_transfer_state(prior, current)
            if bank_conflict:
                conflicts.append({
                    "order_number": row["order_number"],
                    "code": bank_conflict,
                    "prior_snapshot_id": prior.get("latest_snapshot_id"),
                    "incoming_snapshot_id": snapshot_id,
                    "bank_transfer_receipt_review_id": prior.get("bank_transfer_receipt_review_id"),
                })
        current_updates.append(current)

    if snapshots:
        await db.mz2_salla_order_snapshots.insert_many(snapshots)
    for current in current_updates:
        await db.mz2_salla_order_evidence.replace_one(
            {"_id": current["_id"], "user_id": owner},
            current,
            upsert=True,
        )

    file_doc = {
        "_id": file_key,
        "id": file_id,
        "user_id": owner,
        "filename": filename,
        "file_hash": file_hash,
        "source": "salla_orders_export",
        "row_count": len(parsed["rows"]),
        "error_count": len(parsed["errors"]),
        "errors": parsed["errors"][:200],
        "counts": counts,
        "conflict_count": len(conflicts),
        "conflicts": conflicts[:200],
        "uploaded_by": actor["id"],
        "uploaded_at": now,
    }
    await db.mz2_salla_order_files.insert_one(file_doc)
    await db.mz2_salla_order_audit.insert_one({
        "user_id": owner,
        "action": "salla_order_file_imported",
        "file_id": file_id,
        "file_hash": file_hash,
        "row_count": len(parsed["rows"]),
        "error_count": len(parsed["errors"]),
        "conflict_count": len(conflicts),
        "actor_id": actor["id"],
        "at": now,
    })

    return {
        "status": "imported",
        "file": {k: v for k, v in file_doc.items() if k != "_id"},
        "items": [{k: v for k, v in row.items() if k != "_id"} for row in snapshots],
        "errors": parsed["errors"],
        "conflicts": conflicts,
    }


def install_salla_order_evidence_routes(router, db, current_user) -> None:
    base = "/accounting-module/order-evidence"

    async def actor_scope(user: dict[str, Any]) -> tuple[dict[str, Any], str]:
        actor = await fresh_accounting_user(db, user)
        require_accounting_permission(actor, "accounting.receivables.post")
        owner = accounting_owner_id(actor)
        if not owner:
            raise HTTPException(403, "accounting_owner_scope_missing")
        return actor, owner

    @router.get(base)
    async def list_order_evidence(
        limit: int = 200,
        status: str | None = None,
        user: dict = Depends(current_user),
    ):
        actor = await fresh_accounting_user(db, user)
        require_accounting_permission(actor, "accounting.home.view")
        owner = accounting_owner_id(actor)
        query: dict[str, Any] = {"user_id": owner}
        if status:
            query["status"] = status
        items = await db.mz2_salla_order_evidence.find(
            query, {"_id": 0}
        ).sort([("updated_source_text", -1), ("order_number", -1)]).limit(min(max(limit, 1), 500)).to_list(500)
        return {"items": items}

    @router.post(base + "/upload")
    async def upload_order_evidence(
        file: UploadFile = File(...),
        user: dict = Depends(current_user),
    ):
        actor, owner = await actor_scope(user)
        content = await read_safe_xlsx_upload(file, max_bytes=MAX_FILE_BYTES)
        try:
            parsed = parse_salla_order_xlsx(content)
        except ValueError as exc:
            raise HTTPException(400, str(exc)) from exc

        # Structural row errors are retained in the file result so a single
        # unknown payment method or business conflict does not hide the other
        # good rows. Duplicate order numbers inside the same source file are
        # not safe to partially import.
        duplicate_errors = [
            err for err in parsed["errors"]
            if err["code"] == "duplicate_order_number_in_file"
        ]
        if duplicate_errors:
            raise HTTPException(422, detail={
                "code": "duplicate_order_number_in_file",
                "errors": duplicate_errors[:50],
            })

        async def commit(scoped):
            return await import_salla_order_evidence(
                scoped,
                owner=owner,
                actor=actor,
                filename=file.filename or "salla-orders.xlsx",
                content=content,
                parsed=parsed,
            )
        return await atomic_owner(db, owner, commit)
