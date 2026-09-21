"""MZ2-native daily financial evidence and bank-statement intake.

Uploading a bank statement never posts a journal.  Rows are immutable source
facts.  A provider receipt is created only when the provider is explicit in
the uploaded row or an accountant confirms a suggested provider; amount-only
matching never decides identity.
"""
from __future__ import annotations

from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
import hashlib
import io
import json
import re
import uuid
from typing import Any, Literal
from zoneinfo import ZoneInfo

from fastapi import Depends, File, Form, HTTPException, UploadFile
from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict, Field

from accounting_atomic import atomic_owner
from accounting_mz2_balances import read_mz2_write_balances
from accounting_module_contract import (
    OPERATION_ID,
    accounting_owner_id,
    require_accounting_permission,
)
from accounting_module_status_routes import fresh_accounting_user
from accounting_receipt_service import create_receipt
from accounting_settlement_routes import (
    _find_bank,
    _verified_binding_bank_id,
)
from accounting_source_files import preserve_original
from excel_upload_security import read_safe_xlsx_upload
from ledger_core import post_txn_group


MAX_FILE_BYTES = 8 * 1024 * 1024
MAX_ROWS = 5000
PROVIDERS = ("salla", "tamara", "tabby", "emkan")
PROVIDER_ALIASES = {
    "salla": ("salla", "سلة", "سله"),
    "tamara": ("tamara", "تمارا"),
    "tabby": ("tabby", "تابي"),
    "emkan": ("emkan", "imkan", "إمكان", "امكان"),
}
RIYADH = ZoneInfo("Asia/Riyadh")
GENERAL_EXPENSE_CATEGORIES = {
    "rent": "إيجارات",
    "office": "مكتبية",
    "maintenance": "صيانة",
    "subscriptions": "اشتراكات",
    "fuel": "وقود",
    "telecom": "اتصالات وإنترنت",
    "bank_fees": "رسوم بنكية",
    "advertising": "إعلانات",
    "operating_supplies": "مشتريات تشغيلية بسيطة",
    "other": "أخرى",
}
RESERVED_EXPENSE_CODES = {
    "salary", "shipping", "inventory", "tamara_fees", "tabby_fees",
    "gateway_fees", "cod_fees",
}

HEADER_ALIASES = {
    "date": {
        "date", "transaction date", "transaction_date", "posting date",
        "value date", "تاريخ", "تاريخ الحركة", "تاريخ العملية", "تاريخ القيد",
    },
    "credit": {
        "credit", "credit amount", "credit_amount", "deposit", "in",
        "دائن", "إيداع", "ايداع", "وارد", "مبلغ دائن",
    },
    "debit": {
        "debit", "debit amount", "debit_amount", "withdrawal", "out",
        "مدين", "سحب", "صادر", "مبلغ مدين",
    },
    "amount": {
        "amount", "transaction amount", "transaction_amount", "المبلغ", "قيمة العملية",
    },
    "direction": {
        "direction", "type", "transaction type", "اتجاه", "نوع الحركة", "نوع العملية",
    },
    "description": {
        "description", "details", "narrative", "memo", "statement",
        "الوصف", "البيان", "التفاصيل", "شرح", "وصف العملية",
    },
    "reference": {
        "reference", "ref", "transaction id", "transaction_id", "reference number",
        "المرجع", "رقم المرجع", "رقم العملية", "مرجع العملية",
    },
    "provider": {
        "provider", "payment provider", "platform", "المزود", "المنصة", "بوابة الدفع",
    },
}

IN_DIRECTIONS = {"in", "credit", "deposit", "inflow", "وارد", "إيداع", "ايداع", "دائن"}
OUT_DIRECTIONS = {"out", "debit", "withdrawal", "outflow", "صادر", "سحب", "مدين"}


def _clean(value: Any) -> str:
    # Preserve numeric zero; it is a valid debit/credit cell value.
    if value is None:
        return ""
    return " ".join(str(value).strip().split())


def _header(value: Any) -> str:
    return _clean(value).lower().replace("-", " ").replace("_", " ")


def _hash(payload: Any) -> str:
    raw = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str)
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()


def _file_hash(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def _money(value: Any, *, allow_zero=False) -> Decimal:
    if value in (None, ""):
        return Decimal("0.00")
    text = _clean(value).replace(",", "")
    try:
        number = Decimal(text)
    except (InvalidOperation, ValueError):
        raise ValueError("movement_amount_invalid") from None
    if not number.is_finite() or number < 0 or (not allow_zero and number == 0):
        raise ValueError("movement_amount_invalid")
    if number != number.quantize(Decimal("0.01")):
        raise ValueError("movement_amount_precision")
    return number.quantize(Decimal("0.01"))


def _date(value: Any) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = _clean(value)
    if not text:
        raise ValueError("movement_date_required")
    for fmt in (
        "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d",
        "%d/%m/%y", "%m/%d/%Y",
    ):
        try:
            return datetime.strptime(text[:10], fmt).date().isoformat()
        except ValueError:
            pass
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00")).date().isoformat()
    except ValueError:
        raise ValueError("movement_date_invalid") from None


def _movement_accounting_at(movement_date: str) -> datetime:
    try:
        day = datetime.strptime(movement_date, "%Y-%m-%d").date()
    except ValueError:
        raise HTTPException(409, "daily_movement_date_invalid") from None
    return datetime(
        day.year, day.month, day.day, 0, 0, tzinfo=RIYADH
    ).astimezone(timezone.utc)


async def _post_cutover_accounting_at(
    db,
    *,
    owner: str,
    movement_date: str,
) -> str:
    settings = await db.settings.find_one(
        {"user_id": owner},
        {"_id": 0, "mezan2_financial_cutover": 1},
    )
    state = (settings or {}).get("mezan2_financial_cutover") or {}
    if state.get("operation_id") != OPERATION_ID or state.get("status") != "active":
        raise HTTPException(409, "mz2_p01_not_active")
    try:
        cutover = datetime.fromisoformat(
            str(state.get("cutover_at") or "").replace("Z", "+00:00")
        )
    except ValueError:
        raise HTTPException(409, "mz2_cutover_invalid") from None
    if cutover.tzinfo is None or cutover.utcoffset() is None:
        raise HTTPException(409, "mz2_cutover_invalid")
    accounting_at = _movement_accounting_at(movement_date)
    if accounting_at <= cutover.astimezone(timezone.utc):
        raise HTTPException(409, detail={
            "code": "outgoing_event_not_after_cutover",
            "cutover_at": cutover.astimezone(timezone.utc).isoformat(),
            "accounting_at": accounting_at.isoformat(),
        })
    if accounting_at > datetime.now(timezone.utc):
        raise HTTPException(409, "outgoing_event_in_future")
    return accounting_at.isoformat()


def _provider(value: Any) -> str | None:
    text = _clean(value).lower()
    if not text:
        return None
    for provider, aliases in PROVIDER_ALIASES.items():
        if text == provider or text in {alias.lower() for alias in aliases}:
            return provider
    raise ValueError("movement_provider_invalid")


def _suggest_provider(text: str) -> str | None:
    normalized = text.lower()
    matches = []
    for provider, aliases in PROVIDER_ALIASES.items():
        if any(alias.lower() in normalized for alias in aliases):
            matches.append(provider)
    return matches[0] if len(set(matches)) == 1 else None


def _column_map(headers: list[Any]) -> dict[str, int]:
    result: dict[str, int] = {}
    normalized = [_header(value) for value in headers]
    for key, aliases in HEADER_ALIASES.items():
        candidates = {alias.lower().replace("-", " ").replace("_", " ") for alias in aliases}
        for index, value in enumerate(normalized):
            if value in candidates:
                if key in result:
                    raise ValueError(f"duplicate_movement_column:{key}")
                result[key] = index
    if "date" not in result:
        raise ValueError("movement_date_column_required")
    if "description" not in result and "reference" not in result:
        raise ValueError("movement_description_or_reference_required")
    has_split = "credit" in result or "debit" in result
    has_amount_direction = "amount" in result and "direction" in result
    if not has_split and not has_amount_direction:
        raise ValueError("movement_amount_columns_required")
    return result


def _find_header(sheet) -> tuple[int, dict[str, int]]:
    for row_index, row in enumerate(sheet.iter_rows(min_row=1, max_row=20, values_only=True), start=1):
        try:
            mapping = _column_map(list(row))
            return row_index, mapping
        except ValueError:
            continue
    raise ValueError("movement_header_not_detected")


def _cell(row: tuple[Any, ...], mapping: dict[str, int], key: str) -> Any:
    index = mapping.get(key)
    return row[index] if index is not None and index < len(row) else None


def parse_daily_movement_xlsx(content: bytes) -> dict[str, Any]:
    workbook = load_workbook(
        io.BytesIO(content),
        read_only=True,
        data_only=True,
        keep_links=False,
    )
    try:
        sheet = workbook.active
        header_row, mapping = _find_header(sheet)
        rows: list[dict[str, Any]] = []
        errors: list[dict[str, Any]] = []
        for row_no, row in enumerate(
            sheet.iter_rows(min_row=header_row + 1, values_only=True),
            start=header_row + 1,
        ):
            if not any(value not in (None, "") for value in row):
                continue
            if len(rows) + len(errors) >= MAX_ROWS:
                raise ValueError("movement_row_limit")
            try:
                movement_date = _date(_cell(row, mapping, "date"))
                description = _clean(_cell(row, mapping, "description"))
                reference = _clean(_cell(row, mapping, "reference")).upper()
                explicit_provider = _provider(_cell(row, mapping, "provider"))

                if "amount" in mapping:
                    amount = _money(_cell(row, mapping, "amount"))
                    direction_raw = _clean(_cell(row, mapping, "direction")).lower()
                    if direction_raw in IN_DIRECTIONS:
                        direction = "in"
                    elif direction_raw in OUT_DIRECTIONS:
                        direction = "out"
                    else:
                        raise ValueError("movement_direction_invalid")
                else:
                    credit = _money(_cell(row, mapping, "credit"), allow_zero=True)
                    debit = _money(_cell(row, mapping, "debit"), allow_zero=True)
                    if credit > 0 and debit == 0:
                        direction, amount = "in", credit
                    elif debit > 0 and credit == 0:
                        direction, amount = "out", debit
                    else:
                        raise ValueError("movement_credit_debit_conflict")

                if not description and not reference:
                    raise ValueError("movement_description_or_reference_required")
                suggestion = _suggest_provider(" ".join(filter(None, [description, reference])))
                facts = {
                    "row_no": row_no,
                    "date": movement_date,
                    "direction": direction,
                    "amount": format(amount, ".2f"),
                    "description": description,
                    "reference": reference or None,
                    "explicit_provider": explicit_provider,
                    "suggested_provider": suggestion,
                }
                facts["row_hash"] = _hash(facts)
                rows.append(facts)
            except ValueError as exc:
                errors.append({"row_no": row_no, "code": str(exc)})
        if not rows and not errors:
            raise ValueError("movement_file_has_no_rows")
        return {
            "header_row": header_row,
            "mapping": mapping,
            "rows": rows,
            "errors": errors,
        }
    finally:
        workbook.close()


async def _actor_scope(db, user: dict[str, Any], permission: str) -> tuple[dict[str, Any], str]:
    actor = await fresh_accounting_user(db, user)
    require_accounting_permission(actor, permission)
    owner = accounting_owner_id(actor)
    if not owner:
        raise HTTPException(403, "accounting_owner_scope_missing")
    return actor, owner


async def _bank_or_409(db, owner: str, bank_account_id: str) -> dict[str, Any]:
    bank = await _find_bank(db, owner, bank_account_id)
    if not bank or bank.get("account_type") not in {"bank", "cash"}:
        raise HTTPException(409, "daily_movement_bank_invalid")
    return bank


async def _provider_is_confirmable(db, owner: str, bank_account_id: str, provider: str) -> bool:
    return (
        provider in PROVIDERS
        and await _verified_binding_bank_id(db, owner, provider) == bank_account_id
    )


async def _create_provider_receipt_from_movement(
    db,
    *,
    owner: str,
    actor: dict[str, Any],
    movement: dict[str, Any],
    provider: str,
) -> dict[str, Any]:
    if movement.get("direction") != "in":
        raise HTTPException(409, "provider_receipt_requires_inflow")
    if not await _provider_is_confirmable(db, owner, movement["bank_account_id"], provider):
        raise HTTPException(409, "provider_bank_binding_mismatch")
    if movement.get("receipt_id"):
        receipt = await db.mz2_bank_receipts.find_one(
            {"user_id": owner, "id": movement["receipt_id"]},
            {"_id": 0},
        )
        if receipt:
            return receipt

    message_parts = [movement.get("description") or "Bank statement movement"]
    if movement.get("reference"):
        message_parts.append("reference: " + movement["reference"])
    message = " | ".join(message_parts)
    receipt = await create_receipt(
        db,
        owner=owner,
        actor=actor,
        provider=provider,
        amount=movement["amount"],
        bank_message=message,
        received_on=movement["movement_date"],
        request_id="daily-movement:" + movement["id"],
    )
    await db.mz2_bank_receipts.update_one(
        {"user_id": owner, "id": receipt["id"]},
        {"$set": {
            "source": "bank_statement_import",
            "source_daily_movement_id": movement["id"],
            "source_file_id": movement["file_id"],
            "source_file_hash": movement["file_hash"],
            "source_row_no": movement["row_no"],
        }},
    )
    await db.mz2_daily_movements.update_one(
        {"_id": movement["_id"], "user_id": owner},
        {"$set": {
            "status": "provider_receipt_created",
            "confirmed_provider": provider,
            "receipt_id": receipt["id"],
            "classified_by": actor["id"],
            "classified_at": datetime.now(timezone.utc).isoformat(),
        }},
    )
    return receipt


async def import_daily_movement_file(
    db,
    *,
    owner: str,
    actor: dict[str, Any],
    bank_account_id: str,
    filename: str,
    content: bytes,
    parsed: dict[str, Any],
) -> dict[str, Any]:
    bank = await _bank_or_409(db, owner, bank_account_id)
    content_hash = _file_hash(content)
    file_key = _hash([owner, bank_account_id, content_hash])
    prior = await db.mz2_daily_movement_files.find_one({"_id": file_key, "user_id": owner})
    if prior:
        items = await db.mz2_daily_movements.find(
            {"user_id": owner, "file_id": prior["id"]},
            {"_id": 0},
        ).sort("row_no", 1).to_list(MAX_ROWS)
        return {
            "status": "duplicate",
            "file": {k: v for k, v in prior.items() if k != "_id"},
            "items": items,
            "errors": prior.get("errors") or [],
        }

    file_id = str(uuid.uuid4())
    now = datetime.now(timezone.utc).isoformat()
    file_doc = {
        "_id": file_key,
        "id": file_id,
        "user_id": owner,
        "filename": filename,
        "file_hash": content_hash,
        "bank_account_id": bank_account_id,
        "bank_account_name": bank.get("name") or "",
        "uploaded_by": actor["id"],
        "uploaded_at": now,
        "row_count": len(parsed["rows"]),
        "error_count": len(parsed["errors"]),
        "errors": parsed["errors"][:200],
        "header_row": parsed["header_row"],
        "mapping": parsed["mapping"],
        "source": "mz2_bank_statement",
    }

    docs = []
    seen_reference_keys: set[str] = set()
    for row in parsed["rows"]:
        reference = row.get("reference")
        if reference:
            identity = _hash([owner, bank_account_id, "reference", reference])
            if identity in seen_reference_keys:
                raise HTTPException(409, detail={
                    "code": "duplicate_bank_reference_in_file",
                    "reference": reference,
                })
            seen_reference_keys.add(identity)
            prior_reference = await db.mz2_daily_movements.find_one(
                {"_id": identity, "user_id": owner},
            )
            if prior_reference:
                if prior_reference.get("source") in {
                    "manual_accountant",
                    "manual_reconciled_bank_statement",
                }:
                    same_economics = (
                        prior_reference.get("bank_account_id") == bank_account_id
                        and prior_reference.get("direction") == row["direction"]
                        and str(prior_reference.get("amount")) == row["amount"]
                        and prior_reference.get("movement_date") == row["date"]
                    )
                    if not same_economics:
                        raise HTTPException(409, detail={
                            "code": "manual_movement_bank_statement_conflict",
                            "reference": reference,
                            "movement_id": prior_reference["id"],
                        })
                    await db.mz2_daily_movements.update_one(
                        {"_id": identity, "user_id": owner},
                        {"$set": {
                            "file_id": file_id,
                            "file_hash": content_hash,
                            "row_no": row["row_no"],
                            "bank_statement_description": row["description"],
                            "source": "manual_reconciled_bank_statement",
                            "bank_statement_confirmed_at": now,
                        }, "$addToSet": {
                            "evidence_sources": "bank_statement_import",
                        }},
                    )
                    continue
                raise HTTPException(409, detail={
                    "code": "bank_reference_already_imported",
                    "reference": reference,
                    "movement_id": prior_reference["id"],
                })
        else:
            identity = _hash([owner, file_id, row["row_no"], row["row_hash"]])

        explicit = row.get("explicit_provider")
        can_auto = bool(
            explicit
            and row["direction"] == "in"
            and await _provider_is_confirmable(db, owner, bank_account_id, explicit)
        )
        suggestion = row.get("suggested_provider")
        suggestion_valid = bool(
            suggestion
            and row["direction"] == "in"
            and await _provider_is_confirmable(db, owner, bank_account_id, suggestion)
        )
        movement = {
            "_id": identity,
            "id": str(uuid.uuid4()),
            "user_id": owner,
            "file_id": file_id,
            "file_hash": content_hash,
            "row_no": row["row_no"],
            "movement_date": row["date"],
            "direction": row["direction"],
            "amount": row["amount"],
            "currency": "SAR",
            "description": row["description"],
            "reference": reference,
            "bank_account_id": bank_account_id,
            "bank_account_name": bank.get("name") or "",
            "explicit_provider": explicit,
            "suggested_provider": suggestion if suggestion_valid else None,
            "status": "pending_provider_receipt" if can_auto else (
                "needs_review" if suggestion_valid else "unclassified"
            ),
            "receipt_id": None,
            "created_at": now,
            "source": "bank_statement_import",
        }
        docs.append(movement)

    await preserve_original(db, owner, file_id, content)
    await db.mz2_daily_movement_files.insert_one(file_doc)
    if docs:
        await db.mz2_daily_movements.insert_many(docs)
    await db.mz2_daily_movement_audit.insert_one({
        "user_id": owner,
        "action": "file_imported",
        "file_id": file_id,
        "file_hash": content_hash,
        "bank_account_id": bank_account_id,
        "row_count": len(docs),
        "actor_id": actor["id"],
        "at": now,
    })

    auto_created = 0
    for movement in docs:
        if movement["status"] == "pending_provider_receipt":
            await _create_provider_receipt_from_movement(
                db,
                owner=owner,
                actor=actor,
                movement=movement,
                provider=movement["explicit_provider"],
            )
            auto_created += 1

    items = await db.mz2_daily_movements.find(
        {"user_id": owner, "file_id": file_id},
        {"_id": 0},
    ).sort("row_no", 1).to_list(MAX_ROWS)
    return {
        "status": "imported",
        "file": {k: v for k, v in file_doc.items() if k != "_id"},
        "items": items,
        "errors": parsed["errors"],
        "provider_receipts_created": auto_created,
        "needs_review": sum(1 for row in items if row["status"] == "needs_review"),
        "unclassified": sum(1 for row in items if row["status"] == "unclassified"),
    }


class ManualIncomingMovementIn(BaseModel):
    model_config = ConfigDict(extra="forbid")
    bank_account_id: str = Field(min_length=1, max_length=200)
    amount: Decimal
    sender_name: str = Field(min_length=1, max_length=300)
    movement_date: str = Field(min_length=10, max_length=40)
    reference: str = Field(default="", max_length=200)
    notes: str = Field(default="", max_length=1000)
    request_id: str = Field(min_length=8, max_length=200)


async def create_manual_incoming_movement(
    db,
    *,
    owner: str,
    actor: dict[str, Any],
    payload: ManualIncomingMovementIn,
) -> dict[str, Any]:
    bank = await _bank_or_409(db, owner, payload.bank_account_id)
    if bank.get("account_type") != "bank":
        raise HTTPException(409, "manual_bank_transfer_requires_bank_account")

    amount = _money(payload.amount)
    movement_date = _date(payload.movement_date)
    sender = _clean(payload.sender_name)
    reference = _clean(payload.reference).upper()
    notes = _clean(payload.notes)
    if not sender:
        raise HTTPException(422, "manual_transfer_sender_required")

    facts = {
        "bank_account_id": payload.bank_account_id,
        "amount": format(amount, ".2f"),
        "movement_date": movement_date,
        "direction": "in",
        "sender_name": sender,
        "reference": reference or None,
        "notes": notes,
    }
    request_key = _hash([owner, "manual_movement_request", payload.request_id])
    prior_request = await db.mz2_manual_movement_requests.find_one(
        {"_id": request_key, "user_id": owner}
    )
    if prior_request:
        if prior_request.get("facts") != facts:
            raise HTTPException(409, "manual_movement_request_conflict")
        existing = await db.mz2_daily_movements.find_one(
            {"user_id": owner, "id": prior_request["movement_id"]},
            {"_id": 0},
        )
        if not existing:
            raise HTTPException(409, "manual_movement_result_requires_recovery")
        return {**existing, "duplicate": True}

    identity = (
        _hash([owner, payload.bank_account_id, "reference", reference])
        if reference
        else _hash([owner, "manual_movement", payload.request_id])
    )
    existing = await db.mz2_daily_movements.find_one(
        {"_id": identity, "user_id": owner}
    )
    if existing:
        existing_facts = {
            "bank_account_id": existing.get("bank_account_id"),
            "amount": str(existing.get("amount")),
            "movement_date": existing.get("movement_date"),
            "direction": existing.get("direction"),
            "sender_name": existing.get("sender_name") or "",
            "reference": existing.get("reference"),
            "notes": existing.get("manual_notes") or "",
        }
        if (
            existing.get("bank_account_id") != facts["bank_account_id"]
            or str(existing.get("amount")) != facts["amount"]
            or existing.get("movement_date") != facts["movement_date"]
            or existing.get("direction") != "in"
        ):
            raise HTTPException(409, "manual_movement_reference_conflict")
        await db.mz2_manual_movement_requests.insert_one({
            "_id": request_key,
            "user_id": owner,
            "movement_id": existing["id"],
            "facts": facts,
            "created_by": actor["id"],
            "created_at": datetime.now(timezone.utc).isoformat(),
        })
        return {**{k: v for k, v in existing.items() if k != "_id"}, "duplicate": True}

    suggestion = _suggest_provider(" ".join(filter(None, [sender, reference, notes])))
    suggestion_valid = bool(
        suggestion
        and await _provider_is_confirmable(
            db, owner, payload.bank_account_id, suggestion
        )
    )
    movement_id = str(uuid.uuid4())
    now = datetime.now(timezone.utc).isoformat()
    description = f"تحويل وارد — {sender}"
    if notes:
        description += " — " + notes
    movement = {
        "_id": identity,
        "id": movement_id,
        "user_id": owner,
        "file_id": None,
        "file_hash": None,
        "row_no": None,
        "movement_date": movement_date,
        "direction": "in",
        "amount": format(amount, ".2f"),
        "currency": "SAR",
        "description": description,
        "reference": reference or None,
        "bank_account_id": payload.bank_account_id,
        "bank_account_name": bank.get("name") or "",
        "sender_name": sender,
        "manual_notes": notes,
        "explicit_provider": None,
        "suggested_provider": suggestion if suggestion_valid else None,
        "status": "needs_review" if suggestion_valid else "unclassified",
        "receipt_id": None,
        "created_at": now,
        "created_by": actor["id"],
        "source": "manual_accountant",
        "evidence_sources": ["manual_accountant"],
        "request_id": payload.request_id,
    }
    await db.mz2_daily_movements.insert_one(movement)
    await db.mz2_manual_movement_requests.insert_one({
        "_id": request_key,
        "user_id": owner,
        "movement_id": movement_id,
        "facts": facts,
        "created_by": actor["id"],
        "created_at": now,
    })
    await db.mz2_daily_movement_audit.insert_one({
        "user_id": owner,
        "action": "manual_incoming_transfer_created",
        "movement_id": movement_id,
        "bank_account_id": payload.bank_account_id,
        "amount": format(amount, ".2f"),
        "sender_name": sender,
        "reference": reference or None,
        "actor_id": actor["id"],
        "at": now,
    })
    return {k: v for k, v in movement.items() if k != "_id"}


class ManualOutgoingMovementIn(BaseModel):
    model_config = ConfigDict(extra="forbid")
    bank_account_id: str = Field(min_length=1, max_length=200)
    amount: Decimal
    payee_name: str = Field(min_length=1, max_length=300)
    movement_date: str = Field(min_length=10, max_length=40)
    reference: str = Field(default="", max_length=200)
    notes: str = Field(default="", max_length=1000)
    request_id: str = Field(min_length=8, max_length=200)


async def create_manual_outgoing_movement(
    db,
    *,
    owner: str,
    actor: dict[str, Any],
    payload: ManualOutgoingMovementIn,
) -> dict[str, Any]:
    bank = await _bank_or_409(db, owner, payload.bank_account_id)
    if bank.get("account_type") != "bank":
        raise HTTPException(409, "manual_bank_transfer_requires_bank_account")

    amount = _money(payload.amount)
    movement_date = _date(payload.movement_date)
    payee = _clean(payload.payee_name)
    reference = _clean(payload.reference).upper()
    notes = _clean(payload.notes)
    if not payee:
        raise HTTPException(422, "manual_transfer_payee_required")

    facts = {
        "bank_account_id": payload.bank_account_id,
        "amount": format(amount, ".2f"),
        "movement_date": movement_date,
        "direction": "out",
        "payee_name": payee,
        "reference": reference or None,
        "notes": notes,
    }
    request_key = _hash([owner, "manual_movement_request", payload.request_id])
    prior_request = await db.mz2_manual_movement_requests.find_one(
        {"_id": request_key, "user_id": owner}
    )
    if prior_request:
        if prior_request.get("facts") != facts:
            raise HTTPException(409, "manual_movement_request_conflict")
        existing = await db.mz2_daily_movements.find_one(
            {"user_id": owner, "id": prior_request["movement_id"]},
            {"_id": 0},
        )
        if not existing:
            raise HTTPException(409, "manual_movement_result_requires_recovery")
        return {**existing, "duplicate": True}

    identity = (
        _hash([owner, payload.bank_account_id, "reference", reference])
        if reference
        else _hash([owner, "manual_movement", payload.request_id])
    )
    existing = await db.mz2_daily_movements.find_one(
        {"_id": identity, "user_id": owner}
    )
    if existing:
        same_economics = (
            existing.get("bank_account_id") == facts["bank_account_id"]
            and str(existing.get("amount")) == facts["amount"]
            and existing.get("movement_date") == facts["movement_date"]
            and existing.get("direction") == "out"
        )
        if not same_economics:
            raise HTTPException(409, "manual_movement_reference_conflict")
        await db.mz2_manual_movement_requests.insert_one({
            "_id": request_key,
            "user_id": owner,
            "movement_id": existing["id"],
            "facts": facts,
            "created_by": actor["id"],
            "created_at": datetime.now(timezone.utc).isoformat(),
        })
        return {**{k: v for k, v in existing.items() if k != "_id"}, "duplicate": True}

    movement_id = str(uuid.uuid4())
    now = datetime.now(timezone.utc).isoformat()
    description = f"تحويل خارج — {payee}"
    if notes:
        description += " — " + notes
    movement = {
        "_id": identity,
        "id": movement_id,
        "user_id": owner,
        "file_id": None,
        "file_hash": None,
        "row_no": None,
        "movement_date": movement_date,
        "direction": "out",
        "amount": format(amount, ".2f"),
        "currency": "SAR",
        "description": description,
        "reference": reference or None,
        "bank_account_id": payload.bank_account_id,
        "bank_account_name": bank.get("name") or "",
        "payee_name": payee,
        "manual_notes": notes,
        "explicit_provider": None,
        "suggested_provider": None,
        "status": "unclassified",
        "receipt_id": None,
        "created_at": now,
        "created_by": actor["id"],
        "source": "manual_accountant",
        "evidence_sources": ["manual_accountant"],
        "request_id": payload.request_id,
    }
    await db.mz2_daily_movements.insert_one(movement)
    await db.mz2_manual_movement_requests.insert_one({
        "_id": request_key,
        "user_id": owner,
        "movement_id": movement_id,
        "facts": facts,
        "created_by": actor["id"],
        "created_at": now,
    })
    await db.mz2_daily_movement_audit.insert_one({
        "user_id": owner,
        "action": "manual_outgoing_transfer_created",
        "movement_id": movement_id,
        "bank_account_id": payload.bank_account_id,
        "amount": format(amount, ".2f"),
        "payee_name": payee,
        "reference": reference or None,
        "actor_id": actor["id"],
        "at": now,
    })
    return {k: v for k, v in movement.items() if k != "_id"}


class OutgoingMovementClassifyIn(BaseModel):
    model_config = ConfigDict(extra="forbid")
    action: Literal["expense", "supplier_payment"]
    expense_category: str = Field(default="", max_length=100)
    supplier_id: str = Field(default="", max_length=200)
    reason: str = Field(min_length=3, max_length=500)


async def _expense_category(db, owner: str, code: str) -> dict[str, str]:
    code = _clean(code).lower()
    if not code:
        raise HTTPException(422, "expense_category_required")
    if code in RESERVED_EXPENSE_CODES:
        raise HTTPException(409, "expense_category_reserved_for_dedicated_workflow")
    row = await db.expense_categories.find_one(
        {"user_id": owner, "code": code, "status": {"$ne": "inactive"}},
        {"_id": 0, "code": 1, "name": 1},
    )
    if row:
        return {"code": code, "name": row.get("name") or code}
    if code in GENERAL_EXPENSE_CATEGORIES:
        return {"code": code, "name": GENERAL_EXPENSE_CATEGORIES[code]}
    raise HTTPException(404, "expense_category_not_found")


async def _supplier(db, owner: str, supplier_id: str) -> dict[str, str]:
    supplier_id = _clean(supplier_id)
    if not supplier_id:
        raise HTTPException(422, "supplier_required")
    row = await db.suppliers.find_one(
        {
            "user_id": owner,
            "id": supplier_id,
            "status": {"$nin": ["inactive", "archived", "deleted"]},
        },
        {"_id": 0, "id": 1, "company_name": 1, "name": 1},
    )
    if not row:
        raise HTTPException(404, "supplier_not_found")
    return {
        "id": supplier_id,
        "name": row.get("company_name") or row.get("name") or supplier_id,
    }


async def classify_outgoing_movement(
    db,
    *,
    owner: str,
    actor: dict[str, Any],
    movement_id: str,
    payload: OutgoingMovementClassifyIn,
) -> dict[str, Any]:
    movement = await db.mz2_daily_movements.find_one(
        {"user_id": owner, "id": movement_id},
    )
    if not movement:
        raise HTTPException(404, "daily_movement_not_found")

    existing_event_id = str(movement.get("accounting_event_id") or "")
    if existing_event_id:
        prior = await db.mz2_outgoing_financial_events.find_one(
            {"_id": existing_event_id, "user_id": owner},
            {"_id": 0},
        )
        if prior and prior.get("kind") == payload.action:
            expected_ref = (
                payload.expense_category.strip().lower()
                if payload.action == "expense"
                else payload.supplier_id.strip()
            )
            if prior.get("classification_ref") == expected_ref:
                return {**prior, "state": "already_posted"}
        raise HTTPException(409, "daily_movement_already_consumed")

    if (
        movement.get("receipt_id")
        or movement.get("confirmed_provider")
        or movement.get("explicit_provider")
        or movement.get("suggested_provider")
    ):
        raise HTTPException(409, "provider_movement_cannot_be_general_outgoing")
    if movement.get("status") != "unclassified":
        raise HTTPException(409, detail={
            "code": "daily_movement_not_ready_for_outgoing",
            "status": movement.get("status"),
        })
    if movement.get("direction") != "out":
        raise HTTPException(409, "outgoing_classification_requires_outflow")

    amount = _money(movement.get("amount"))
    bank_id = str(movement.get("bank_account_id") or "").strip()
    if not bank_id:
        raise HTTPException(409, "daily_movement_bank_missing")
    accounting_at = await _post_cutover_accounting_at(
        db, owner=owner, movement_date=movement["movement_date"]
    )

    category = None
    supplier = None
    required = [("bank", bank_id, "main")]
    if payload.action == "expense":
        if payload.supplier_id.strip():
            raise HTTPException(422, "expense_cannot_include_supplier")
        category = await _expense_category(db, owner, payload.expense_category)
        classification_ref = category["code"]
    else:
        if payload.expense_category.strip():
            raise HTTPException(422, "supplier_payment_cannot_include_expense_category")
        supplier = await _supplier(db, owner, payload.supplier_id)
        classification_ref = supplier["id"]
        required.append(("supplier", supplier["id"], "payable"))

    balances = await read_mz2_write_balances(
        db,
        owner=owner,
        required_accounts=required,
    )
    bank_balance = balances.net_balance(
        entity_type="bank", entity_id=bank_id, sub_account="main"
    )
    if amount > bank_balance:
        raise HTTPException(409, detail={
            "code": "insufficient_mz2_bank_balance",
            "available": format(max(bank_balance, Decimal(0)), ".2f"),
            "required": format(amount, ".2f"),
        })

    if payload.action == "supplier_payment":
        payable = max(
            -balances.net_balance(
                entity_type="supplier",
                entity_id=supplier["id"],
                sub_account="payable",
            ),
            Decimal(0),
        )
        if payable <= 0:
            raise HTTPException(409, "supplier_has_no_mz2_payable")
        if amount > payable:
            raise HTTPException(409, detail={
                "code": "supplier_payment_exceeds_payable",
                "payable": format(payable, ".2f"),
                "payment": format(amount, ".2f"),
            })

    event_id = _hash([owner, "outgoing_movement", movement_id])
    economic = {
        "movement_id": movement_id,
        "kind": payload.action,
        "classification_ref": classification_ref,
        "amount": format(amount, ".2f"),
        "accounting_at": accounting_at,
        "bank_account_id": bank_id,
    }
    economic_hash = _hash(economic)
    prior_event = await db.mz2_outgoing_financial_events.find_one(
        {"_id": event_id, "user_id": owner},
    )
    if prior_event:
        if prior_event.get("economic_hash") != economic_hash:
            raise HTTPException(409, "outgoing_movement_accounting_conflict")
        return {**{k: v for k, v in prior_event.items() if k != "_id"}, "state": "already_posted"}

    reason = payload.reason.strip()
    metadata = {
        "operation_id": OPERATION_ID,
        "source": "accounting_daily_outgoing_p01",
        "outgoing_event_id": event_id,
        "daily_movement_id": movement_id,
        "daily_movement_file_id": movement.get("file_id"),
        "bank_reference": movement.get("reference"),
        "accounting_at": accounting_at,
        "reason": reason,
        "classification_ref": classification_ref,
    }
    if payload.action == "expense":
        metadata.update({
            "expense_category": category["code"],
            "expense_category_name": category["name"],
        })
        entries = [
            {
                "entity_type": "expense",
                "entity_id": category["code"],
                "sub_account": None,
                "side": "debit",
                "amount": float(amount),
                "entry_type": "expense_record",
            },
            {
                "entity_type": "bank",
                "entity_id": bank_id,
                "sub_account": "main",
                "side": "credit",
                "amount": float(amount),
                "entry_type": "expense_record",
            },
        ]
        txn_type = "mz2_general_expense"
        note = f"{category['name']} — {reason}"
    else:
        metadata.update({
            "supplier_id": supplier["id"],
            "supplier_name": supplier["name"],
        })
        entries = [
            {
                "entity_type": "supplier",
                "entity_id": supplier["id"],
                "sub_account": "payable",
                "side": "debit",
                "amount": float(amount),
                "entry_type": "supplier_payment",
            },
            {
                "entity_type": "bank",
                "entity_id": bank_id,
                "sub_account": "main",
                "side": "credit",
                "amount": float(amount),
                "entry_type": "supplier_payment",
            },
        ]
        txn_type = "mz2_supplier_payment"
        note = f"سداد مورد {supplier['name']} — {reason}"

    result = await post_txn_group(
        db,
        user_id=owner,
        actor_id=actor["id"],
        actor_name=actor.get("name") or actor.get("email") or actor["id"],
        txn_type=txn_type,
        notes=note,
        metadata=metadata,
        entries=entries,
    )
    now = datetime.now(timezone.utc).isoformat()
    event = {
        "_id": event_id,
        "id": event_id,
        "user_id": owner,
        "kind": payload.action,
        "classification_ref": classification_ref,
        "movement_id": movement_id,
        "amount": format(amount, ".2f"),
        "accounting_at": accounting_at,
        "bank_account_id": bank_id,
        "bank_reference": movement.get("reference"),
        "economic_hash": economic_hash,
        "txn_group_id": result["txn_group_id"],
        "status": "posted",
        "reason": reason,
        "posted_by": actor["id"],
        "posted_at": now,
        **(
            {
                "expense_category": category["code"],
                "expense_category_name": category["name"],
            }
            if payload.action == "expense"
            else {
                "supplier_id": supplier["id"],
                "supplier_name": supplier["name"],
            }
        ),
    }
    await db.mz2_outgoing_financial_events.insert_one(event)
    changed = await db.mz2_daily_movements.update_one(
        {
            "_id": movement["_id"],
            "user_id": owner,
            "accounting_event_id": {"$in": [None, ""]},
        },
        {"$set": {
            "status": "accounting_posted",
            "accounting_event_id": event_id,
            "accounting_action": payload.action,
            "accounting_txn_group_id": result["txn_group_id"],
            "accounting_reason": reason,
            "accounting_classification_ref": classification_ref,
            "consumed_at": now,
            "consumed_by": actor["id"],
        }},
    )
    if changed.modified_count != 1:
        raise HTTPException(409, "daily_movement_concurrent_consumption")
    await db.mz2_daily_movement_audit.insert_one({
        "user_id": owner,
        "action": "outgoing_classified",
        "event_id": event_id,
        "movement_id": movement_id,
        "kind": payload.action,
        "classification_ref": classification_ref,
        "txn_group_id": result["txn_group_id"],
        "actor_id": actor["id"],
        "at": now,
        "reason": reason,
    })
    return {k: v for k, v in event.items() if k != "_id"}


class ProviderConfirmIn(BaseModel):
    model_config = ConfigDict(extra="forbid")
    provider: Literal["salla", "tamara", "tabby", "emkan"]
    reason: str = Field(min_length=3, max_length=500)


async def confirm_daily_movement_provider(
    db,
    *,
    owner: str,
    actor: dict[str, Any],
    movement_id: str,
    payload: ProviderConfirmIn,
) -> dict[str, Any]:
    movement = await db.mz2_daily_movements.find_one(
        {"user_id": owner, "id": movement_id},
    )
    if not movement:
        raise HTTPException(404, "daily_movement_not_found")
    if movement.get("receipt_id"):
        return {k: v for k, v in movement.items() if k != "_id"}
    if movement.get("status") not in {"needs_review", "unclassified", "pending_provider_receipt"}:
        raise HTTPException(409, "daily_movement_not_classifiable")
    receipt = await _create_provider_receipt_from_movement(
        db,
        owner=owner,
        actor=actor,
        movement=movement,
        provider=payload.provider,
    )
    await db.mz2_daily_movements.update_one(
        {"user_id": owner, "id": movement_id},
        {"$set": {
            "classification_reason": payload.reason.strip(),
            "confirmed_provider": payload.provider,
        }},
    )
    await db.mz2_daily_movement_audit.insert_one({
        "user_id": owner,
        "action": "provider_confirmed",
        "movement_id": movement_id,
        "provider": payload.provider,
        "reason": payload.reason.strip(),
        "actor_id": actor["id"],
        "at": datetime.now(timezone.utc).isoformat(),
    })
    current = await db.mz2_daily_movements.find_one(
        {"user_id": owner, "id": movement_id},
        {"_id": 0},
    )
    return {**current, "receipt": receipt}


async def daily_movement_context(db, owner: str) -> dict[str, Any]:
    banks = await db.accounts.find(
        {
            "user_id": owner,
            "status": {"$ne": "hidden"},
            "account_type": {"$in": ["bank", "cash"]},
        },
        {"_id": 0, "id": 1, "name": 1, "account_type": 1},
    ).sort("name", 1).to_list(200)
    bindings = {}
    for provider in PROVIDERS:
        bank_id = await _verified_binding_bank_id(db, owner, provider)
        bindings[provider] = bank_id

    stored_categories = await db.expense_categories.find(
        {
            "user_id": owner,
            "status": {"$ne": "inactive"},
            "code": {"$nin": sorted(RESERVED_EXPENSE_CODES)},
        },
        {"_id": 0, "code": 1, "name": 1},
    ).sort("name", 1).to_list(500)
    categories = {
        code: {"code": code, "name": name, "source": "mz2_builtin"}
        for code, name in GENERAL_EXPENSE_CATEGORIES.items()
    }
    for row in stored_categories:
        code = _clean(row.get("code")).lower()
        if code and code not in RESERVED_EXPENSE_CODES:
            categories[code] = {
                "code": code,
                "name": row.get("name") or code,
                "source": "stored",
            }

    suppliers = await db.suppliers.find(
        {
            "user_id": owner,
            "status": {"$nin": ["inactive", "archived", "deleted"]},
        },
        {"_id": 0, "id": 1, "company_name": 1, "name": 1, "status": 1},
    ).sort("company_name", 1).to_list(1000)
    return {
        "banks": banks,
        "provider_bank_bindings": bindings,
        "supported_providers": list(PROVIDERS),
        "expense_categories": sorted(categories.values(), key=lambda row: row["name"]),
        "suppliers": [
            {
                "id": row["id"],
                "name": row.get("company_name") or row.get("name") or row["id"],
                "status": row.get("status") or "active",
            }
            for row in suppliers
            if row.get("id")
        ],
        "template_columns": [
            "date", "credit", "debit", "description", "reference", "provider"
        ],
        "accounting_effect": "none_until_classified_and_matched",
    }


def install_daily_movement_routes(router, db, current_user) -> None:
    base = "/accounting-module/daily-movements"

    @router.get(base + "/context")
    async def context(user: dict = Depends(current_user)):
        actor, owner = await _actor_scope(db, user, "accounting.movements.view")
        del actor
        return await daily_movement_context(db, owner)

    @router.get(base)
    async def list_movements(
        limit: int = 200,
        status: str | None = None,
        user: dict = Depends(current_user),
    ):
        _, owner = await _actor_scope(db, user, "accounting.movements.view")
        query: dict[str, Any] = {"user_id": owner}
        if status:
            query["status"] = status
        rows = await db.mz2_daily_movements.find(
            query, {"_id": 0}
        ).sort([("movement_date", -1), ("created_at", -1)]).limit(min(max(limit, 1), 500)).to_list(500)
        return {"items": rows}

    @router.post(base + "/manual-incoming")
    async def manual_incoming(
        payload: ManualIncomingMovementIn,
        user: dict = Depends(current_user),
    ):
        actor, owner = await _actor_scope(db, user, "accounting.movements.import")

        async def save(scoped):
            return await create_manual_incoming_movement(
                scoped,
                owner=owner,
                actor=actor,
                payload=payload,
            )
        return await atomic_owner(db, owner, save)

    @router.post(base + "/manual-outgoing")
    async def manual_outgoing(
        payload: ManualOutgoingMovementIn,
        user: dict = Depends(current_user),
    ):
        actor, owner = await _actor_scope(db, user, "accounting.movements.import")

        async def save(scoped):
            return await create_manual_outgoing_movement(
                scoped,
                owner=owner,
                actor=actor,
                payload=payload,
            )
        return await atomic_owner(db, owner, save)

    @router.post(base + "/{movement_id}/classify-outgoing")
    async def classify_outgoing(
        movement_id: str,
        payload: OutgoingMovementClassifyIn,
        user: dict = Depends(current_user),
    ):
        actor, owner = await _actor_scope(db, user, "accounting.movements.view")
        require_accounting_permission(
            actor,
            "accounting.journals.manual_create"
            if payload.action == "expense"
            else "accounting.settlements.post",
        )

        async def post(scoped):
            return await classify_outgoing_movement(
                scoped,
                owner=owner,
                actor=actor,
                movement_id=movement_id,
                payload=payload,
            )
        return await atomic_owner(db, owner, post)

    @router.post(base + "/upload")
    async def upload(
        file: UploadFile = File(...),
        bank_account_id: str = Form(...),
        user: dict = Depends(current_user),
    ):
        actor, owner = await _actor_scope(db, user, "accounting.movements.import")
        content = await read_safe_xlsx_upload(file, max_bytes=MAX_FILE_BYTES)
        try:
            parsed = parse_daily_movement_xlsx(content)
        except ValueError as exc:
            raise HTTPException(400, str(exc)) from exc
        if parsed["errors"]:
            raise HTTPException(422, detail={
                "code": "daily_movement_rows_invalid",
                "error_count": len(parsed["errors"]),
                "errors": parsed["errors"][:50],
                "message": "لم يُستورد الملف لأن بعض صفوف كشف البنك غير صالحة",
            })

        async def save(scoped):
            return await import_daily_movement_file(
                scoped,
                owner=owner,
                actor=actor,
                bank_account_id=bank_account_id,
                filename=file.filename or "bank-statement.xlsx",
                content=content,
                parsed=parsed,
            )
        return await atomic_owner(db, owner, save)

    @router.post(base + "/{movement_id}/confirm-provider")
    async def confirm_provider(
        movement_id: str,
        payload: ProviderConfirmIn,
        user: dict = Depends(current_user),
    ):
        actor, owner = await _actor_scope(db, user, "accounting.movements.import")
        return await confirm_daily_movement_provider(
            db,
            owner=owner,
            actor=actor,
            movement_id=movement_id,
            payload=payload,
        )
